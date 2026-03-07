import docx

def inspect_doc(path):
    print(f"\nInspecting FOOTERS: {path}")
    doc = docx.Document(path)
    for i, sec in enumerate(doc.sections):
        print(f"  Section {i}:")
        for f_name, f_obj in [('footer', sec.footer), ('1st_footer', sec.first_page_footer)]:
            print(f"    {f_name}: {len(f_obj.paragraphs)} paras")
            for j, p in enumerate(f_obj.paragraphs):
                # check runs
                runs_info = [len(r._element) for r in p.runs]
                print(f"      P{j} text='{p.text}' runs={len(p.runs)} XML_children_per_run={runs_info}")

inspect_doc('assets/PORTADA DE ESTADÍA EDITABLE 2026.docx')
