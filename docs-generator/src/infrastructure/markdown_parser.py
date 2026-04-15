"""
File: infrastructure/markdown_parser.py
Purpose: Parses markdown files and injects content into docx preserving formats.
All Rights Reserved Arodi Emmanuel
"""
import os
import re
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from src.domain.constants import FONT_NAME
from src.infrastructure.docx_styles_adapter import _ensure_style

_ASSETS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', 'content', 'assets'))

_FIGURE_MAP = {
    'fig_access_evolution': 'fig_access_evolution.png',
    'fig_problem_triad':    'fig_problem_triad.png',
    'fig_ddd_layers':       'fig_ddd_layers.png',
    'fig_clean_rings':      'fig_clean_rings.png',
    'fig_solid_pillars':    'fig_solid_pillars.png',
    'fig_vdom_diff':        'fig_vdom_diff.png',
    'fig_arch':     'fig_arch.png',
    'fig_erd':      'fig_erd.png',
    'fig_sequence': 'fig_sequence.png',
    'fig_stack':    'fig_stack.png',
    'fig_pyramid':  'fig_pyramid.png',
    'fig_access':   'fig_access.png',
    'fig_dashboard':'fig_dashboard.png',
    'fig_metrics':  'fig_metrics.png',
    'fig_hardware_admin': (
        'fig_admin_hardware_view'
        '(check_finger_print_status_only.png'),
    'fig_theme_change': (
        'fig_cambio_de_tema(otro_tema_ha_sido'
        '_seleccionado_y_la_interfaz_ha_'
        'cambiado).png'),
    'fig_carrousel': 'fig_carrousel.png',
    'fig_carrousel_config': (
        'fig_carrousel_configuration.png'),
    'fig_control_tokens': (
        'fig_control_tokes(languages, themes,,'
        ' font-size,spacing).png'),
    'fig_user_mgmt': (
        'fig_gestión_usuarios.png'),
    'fig_db_backup': (
        'fig_respaldo_de_db_automatico_y_'
        'restauracion_exportacion_de_datos_'
        'en_excel_y_db.png'),
    'fig_membership_sale': (
        'fig_venta_membresía.png'),
    'fig_ci_pipeline': 'fig_ci_pipeline.png',
    'fig_sdk_flow': 'fig_sdk_flow.png',
    'fig_component_tree': 'fig_component_tree.png',
    'fig_psp_levels': 'fig_psp_levels.png',
}

def _add_formatted_runs(p, text):
    """
    Parses a string for **bold** and *italic* markdown and appends
    corresponding runs to the paragraph `p`.
    """
    # Simple regex to tokenize text into formats.
    # It looks for **text**, *text*, or regular text.
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    
    for token in tokens:
        if not token:
            continue
        run = p.add_run()
        run.font.name = FONT_NAME
        run.font.size = Pt(12)
        
        if token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run.text = token[1:-1]
            run.italic = True
        else:
            run.text = token

def _add_body(doc, text):
    p = doc.add_paragraph()
    _add_formatted_runs(p, text)
    fmt = p.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0)
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(0)
    return p


def _add_caption(doc, text):
    """Render figure caption: centered, 10pt italic, small spacing."""
    # Strip surrounding *...*
    raw = re.sub(r'^\*|\*$', '', text.strip())
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.space_before = Pt(2)
    fmt.space_after = Pt(8)
    fmt.line_spacing = 1.15
    run = p.add_run(raw)
    run.font.name = FONT_NAME
    run.font.size = Pt(10)
    run.font.italic = True
    return p

def _add_bullet(doc, text):
    _ensure_style(doc, 'List Bullet')
    p = doc.add_paragraph(style='List Bullet')
    _add_formatted_runs(p, text)
    fmt = p.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0)
    fmt.line_spacing = 1.5
    return p

def _add_numbered(doc, text):
    _ensure_style(doc, 'List Number')
    p = doc.add_paragraph(style='List Number')
    _add_formatted_runs(p, text)
    fmt = p.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0)
    fmt.line_spacing = 1.5
    return p

def _add_image(doc, fig_key):
    filename = _FIGURE_MAP.get(fig_key)
    if not filename:
        return
    path = os.path.join(_ASSETS_DIR, filename)
    if not os.path.exists(path):
        print(f'[WARN] Image not found: {path}')
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(5.8))

def _add_code_block(doc, text):
    text = re.sub(r'^```\w*\n', '', text)
    text = re.sub(r'\n```$', '', text)
    p = doc.add_paragraph(text)
    fmt = p.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt.left_indent = Cm(1)
    fmt.line_spacing = 1.0
    for run in p.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(40, 40, 40)
    return p

def _collapse_list_items(block, marker_re):
    """
    Rejoin multi-line list items into single logical items.

    Markdown files are line-wrapped at ~60 chars. A bullet like:
        - **Diseñar** un núcleo de dominio desacoplado utilizando
        Domain-Driven Design para aislar las reglas de negocio de
        la infraestructura tecnológica.

    must become ONE item, not three. A continuation line is any
    line that does NOT start with a new list marker.
    """
    items = []
    current = None
    for line in block.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if re.match(marker_re, stripped):
            if current is not None:
                items.append(current)
            current = re.sub(marker_re, '', stripped, count=1).strip()
        elif current is not None:
            current += ' ' + stripped
        else:
            current = stripped
    if current is not None:
        items.append(current)
    return items


def parse_markdown_blocks(doc, content):
    content = re.sub(r'/\*.*?\*/', '', content, flags=re.DOTALL)
    content = re.sub(
        r'^(#+ .+)(\n)(?=[^\n])',
        r'\1\n\n', content, flags=re.MULTILINE
    )

    blocks = [
        b.strip()
        for b in content.strip().split('\n\n')
        if b.strip()
    ]
    for block in blocks:
        if block.startswith('```'):
            _add_code_block(doc, block)
            continue
        img_match = re.match(r'^!!(fig_\w+)!!$', block.strip())
        # Detect figure captions: *Figura N.N: ...*
        caption_match = re.match(
            r'^\*Figura\s+\d+[\.,]\d+.*\*$',
            block.strip(), re.DOTALL
        )
        if img_match:
            _add_image(doc, img_match.group(1))
        elif caption_match:
            raw = ' '.join(block.splitlines()).strip()
            _add_caption(doc, raw)
        elif block.startswith('# '):
            doc.add_heading(block[2:].strip(), level=1)
        elif block.startswith('## '):
            doc.add_heading(block[3:].strip(), level=2)
        elif block.startswith('### '):
            doc.add_heading(block[4:].strip(), level=3)
        elif re.match(r'^[-*] ', block):
            for item in _collapse_list_items(block, r'^[-*]\s+'):
                _add_bullet(doc, item)
        elif re.match(r'^\d+\.', block):
            for item in _collapse_list_items(block, r'^\d+\.\s*'):
                _add_numbered(doc, item)
        else:
            lines = [ln.strip() for ln in block.splitlines()]
            text = ' '.join(ln for ln in lines if ln)
            # Collapse any double spaces introduced by line-wrap stripping
            text = re.sub(r'  +', ' ', text)
            _add_body(doc, text)
