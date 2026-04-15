"""
File: infrastructure/docx_styles_adapter.py
Purpose: Configures native MS Word styles and margins for the document.
All Rights Reserved Arodi Emmanuel
"""
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from src.domain.constants import MARGIN_LEFT, MARGIN_RIGHT, MARGIN_TOP, MARGIN_BOTTOM, FONT_NAME

def configure_margins(doc):
    for section in doc.sections:
        section.left_margin = MARGIN_LEFT
        section.right_margin = MARGIN_RIGHT
        section.bottom_margin = MARGIN_BOTTOM
        section.top_margin = MARGIN_TOP

def _ensure_style(doc, name, style_type=1): # 1 is WD_STYLE_TYPE.PARAGRAPH
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, style_type)

def configure_normal_style(doc):
    style = _ensure_style(doc, 'Normal')
    font = style.font
    font.name = FONT_NAME
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    fmt = style.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)

def _set_outline_level(style, level):
    """
    Inject <w:outlineLvl w:val='level'/> into style's pPr so that
    Word's TOC field can detect the heading. Without this tag,
    TOC \\o returns 'No entries found'.
    """
    pPr = style.element.get_or_add_pPr()
    outlineLvl = pPr.find(qn('w:outlineLvl'))
    if outlineLvl is None:
        outlineLvl = OxmlElement('w:outlineLvl')
        pPr.append(outlineLvl)
    outlineLvl.set(qn('w:val'), str(level))

def configure_heading_styles(doc):
    all_headings = {
        'Title':     (16, Pt(0), Pt(0), WD_ALIGN_PARAGRAPH.CENTER, True, None),
        'Heading 1': (16, Pt(0), Pt(0), WD_ALIGN_PARAGRAPH.CENTER, True,  0),
        'Heading 2': (14, Pt(0), Pt(0), WD_ALIGN_PARAGRAPH.LEFT,   False, 1),
        'Heading 3': (12, Pt(0), Pt(0), WD_ALIGN_PARAGRAPH.LEFT,   False, 2),
    }
    for name, (size, before, after, align, bold, outline) in all_headings.items():
        style = _ensure_style(doc, name)
        font = style.font
        font.name = FONT_NAME
        font.size = Pt(size)
        font.bold = bold
        font.underline = False
        font.color.rgb = RGBColor(0, 0, 0)
        fmt = style.paragraph_format
        fmt.alignment = align
        fmt.space_before = before
        fmt.space_after = after
        fmt.line_spacing = 1.5
        if outline is not None:
            _set_outline_level(style, outline)

def _write_page_number(footer_obj):
    """Write PAGE field bottom-right into the given footer object."""
    paragraph = (
        footer_obj.paragraphs[0]
        if footer_obj.paragraphs
        else footer_obj.add_paragraph()
    )
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run.font.name = FONT_NAME
    run.font.size = Pt(10)

def set_page_numbers(doc):
    """Apply PAGE footer to every section except cover."""
    for i, section in enumerate(doc.sections):
        if i == 0:
            continue
        section.footer.is_linked_to_previous = False
        _write_page_number(section.footer)
        section.first_page_footer.is_linked_to_previous = False
        _write_page_number(section.first_page_footer)


def setup_document_styles(doc):
    configure_margins(doc)
    configure_normal_style(doc)
    configure_heading_styles(doc)
    set_page_numbers(doc)
