import docx
import os
import re

def create_thesis_docx():
    doc = docx.Document()
    
    # Optional: Set styles or margins if needed
    
    # Title - Capítulo 1
    doc.add_heading('CAPÍTULO 1. INTRODUCCIÓN Y GENERALIDADES', 0)
    
    chapter_1_dir = 'docs-generator/capitulo_1'
    files = [
        '1.1_introduccion.md',
        '1.2_antecedentes.md',
        '1.3_planteamiento_problema.md',
        '1.4_objetivos.md',
        '1.5_justificación.md',
        '1.6_alcances_limitaciones.md'
    ]
    
    for filename in files:
        filepath = os.path.join(chapter_1_dir, filename)
        if not os.path.exists(filepath):
            continue
            
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
            
        # Remove comments /* ... */
        content = re.sub(r'/\*.*?\*/', '', content, flags=re.DOTALL)
        
        lines = content.strip().split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif any(line.startswith(f"{i}. ") for i in range(1, 10)):
                doc.add_paragraph(line[3:], style='List Number')
            else:
                doc.add_paragraph(line)
        
        # Add a page break or spacing between sections?
        # doc.add_page_break()
        doc.add_paragraph() # Just a space for now
        
    output_path = 'Tesis_Poetry_Borrador.docx'
    doc.save(output_path)
    print(f"Document saved to: {output_path}")

if __name__ == "__main__":
    create_thesis_docx()
